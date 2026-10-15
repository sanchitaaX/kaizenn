"""Read the text out of a project brief the user attaches in the chat (.txt, .md, .pdf, .docx).

Uploaded files are only ever parsed for text. Nothing in them is executed, imported or followed;
the text is handed back as a plain string that the user can still edit before sending.
"""

from __future__ import annotations

import io
import logging
import posixpath
import zipfile

log = logging.getLogger(__name__)

MAX_BYTES = 2 * 1024 * 1024
MAX_TEXT_CHARS = 20000          # same cap as the brief typed in the chat
MAX_PDF_PAGES = 60
MAX_UNZIPPED_BYTES = 25 * 1024 * 1024   # refuse "zip bombs" disguised as .docx
ALLOWED = (".txt", ".md", ".pdf", ".docx")
# A repository brief may also carry code and config the owner wants the agent to read (still only as text).
REPO_ALLOWED = ALLOWED + (".rst", ".json", ".yaml", ".yml", ".toml", ".ini", ".csv", ".sql", ".py", ".js", ".ts",
                          ".tsx", ".jsx", ".html")

WRONG_TYPE = "Attach a .txt, .md, .pdf or .docx file."
REPO_WRONG_TYPE = "Attach a text, Markdown, PDF, Word, code or config file (.txt .md .pdf .docx .json .yaml .py .js ...)."
TOO_BIG = "That file is larger than 2 MB. Attach a smaller one or paste the text instead."
EMPTY = "I could not find any text in that file."


class BriefFileError(ValueError):
    """A plain-language reason the file cannot be used. `status` is the HTTP status to send."""

    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.status = status


def clean_filename(name: str | None) -> str:
    base = posixpath.basename((name or "").replace("\\", "/")).strip()
    base = "".join(ch for ch in base if ch.isprintable())
    return base[:120] or "brief"


def extension(name: str) -> str:
    return posixpath.splitext(name.lower())[1]


def _from_text(data: bytes) -> str:
    if b"\x00" in data:
        raise BriefFileError("That file does not look like plain text.")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise BriefFileError("I can only read text files saved as UTF-8.")


def _from_pdf(data: bytes) -> str:
    if not data.lstrip()[:5].startswith(b"%PDF"):
        raise BriefFileError("That file is not a real PDF.")
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise BriefFileError("That PDF is password protected. Paste the text instead.")
        pages = reader.pages[:MAX_PDF_PAGES]
        return "\n\n".join((page.extract_text() or "") for page in pages)
    except BriefFileError:
        raise
    except Exception:
        log.warning("Could not read a PDF brief", exc_info=True)
        raise BriefFileError("I could not read that PDF. Paste the text instead.")


def _from_docx(data: bytes) -> str:
    if data[:2] != b"PK":
        raise BriefFileError("That file is not a real Word document.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED_BYTES:
                raise BriefFileError("That document is too large to read.")
    except zipfile.BadZipFile:
        raise BriefFileError("That file is not a real Word document.")
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(lines)
    except Exception:
        log.warning("Could not read a Word brief", exc_info=True)
        raise BriefFileError("I could not read that document. Paste the text instead.")


def extract_text(filename: str, data: bytes, for_repo: bool = False) -> dict:
    """Returns {filename, text}. Raises BriefFileError with a plain message. for_repo also accepts code and config."""
    name = clean_filename(filename)
    ext = extension(name)
    if ext not in (REPO_ALLOWED if for_repo else ALLOWED):
        raise BriefFileError(REPO_WRONG_TYPE if for_repo else WRONG_TYPE)
    if len(data) > MAX_BYTES:
        raise BriefFileError(TOO_BIG, status=413)
    if ext == ".pdf":
        text = _from_pdf(data)
    elif ext == ".docx":
        text = _from_docx(data)
    else:
        text = _from_text(data)
    text = "\n".join(line.rstrip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"))
    text = "".join(ch for ch in text if ch in "\n\t" or ch.isprintable()).strip()
    if not text:
        raise BriefFileError(EMPTY)
    return {"filename": name, "text": text[:MAX_TEXT_CHARS]}
